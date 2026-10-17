"""
Tests for history/api.py + history/export.py (Phase 4 + 5).  Run:  python3 test_api_export.py
Needs: fastapi, httpx (TestClient), openpyxl.  No ROS required.
"""
from __future__ import annotations

import io
import math
import os
import shutil
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, "../../..")))

from fastapi import FastAPI                       # noqa: E402
from fastapi.staticfiles import StaticFiles       # noqa: E402
from fastapi.testclient import TestClient         # noqa: E402
from openpyxl import load_workbook                # noqa: E402

from amr_web.history import api, db, events, runs  # noqa: E402

_FAILURES = []


def expect(cond, msg):
    if not cond:
        _FAILURES.append(msg)
        print(f"  FAIL: {msg}")
    else:
        print(f"  ok:   {msg}")


def flush(timeout=3.0):
    """Wait until the writer thread has drained its queue."""
    end = time.time() + timeout
    while time.time() < end and not db._queue.empty():
        time.sleep(0.02)
    time.sleep(0.15)


def make_run(goal, start, end, result, samples, scenario=None, waypoints=None, mtype="waypoints"):
    if waypoints:
        s = runs.begin_waypoint_run(waypoints, start, "map_01", "route_a", mtype)
    else:
        s = runs.begin_goal_run(goal, start, "map_01", "route_a")
    rid = s.run_id
    t0 = time.time()
    for i, (x, y) in enumerate(samples):
        s.sample_pose(t0 + i * 0.2, x, y, 0.0)
    if waypoints:
        runs.on_waypoint_progress(1, {"x": waypoints[0]["x"] + 0.01, "y": waypoints[0]["y"], "yaw": 0.0})
    runs.finish_active_run(result, f"{result} test", end)
    if scenario:
        db.update_run_meta(rid, scenario=scenario)
    return rid


def build(tmp):
    db.shutdown()
    runs.set_active(None)
    db.init(db_path=os.path.join(tmp, "api.db"))
    now = time.time()
    # events
    e1 = events.record_command({"command": "start_slam"}, client_ip="10.0.0.5")
    e2 = events.record_command({"command": "nav_goal", "pose": {"x": 1, "y": 2, "yaw": 0}}, client_ip="10.0.0.6")
    events.update_command_result(e1, True, "SLAM started")
    events.update_command_result(e2, False, "Nav2 is not running")
    events.record_system_event("SYSTEM", "device_disconnected", level="warn", detail="device=lidar")
    # runs
    line = [(i * 0.1, 0.0) for i in range(11)]                    # 1 m
    r1 = make_run({"x": 1.0, "y": 0.0, "yaw": 0.0}, {"x": 0.0, "y": 0.0, "yaw": 0.0},
                  {"x": 1.002, "y": -0.004, "yaw": 0.01}, "SUCCEEDED", line, scenario="A to B")
    r2 = make_run({"x": 2.0, "y": 0.0, "yaw": 0.0}, {"x": 0.0, "y": 0.0, "yaw": 0.0},
                  {"x": 1.0, "y": 0.0, "yaw": 0.0}, "CANCELED", line[:6], scenario="A to B")
    r3 = make_run(None, {"x": 0.0, "y": 0.0, "yaw": 0.0}, {"x": 2.0, "y": 0.0, "yaw": 0.0}, "SUCCEEDED",
                  [(i * 0.2, 0.0) for i in range(11)], scenario="Via points",
                  waypoints=[{"x": 1.0, "y": 0.0, "yaw": 0.0, "name": "P1"},
                             {"x": 2.0, "y": 0.0, "yaw": 0.0, "name": "P2"}])
    flush()
    return now, (r1, r2, r3)


def xl(resp):
    return load_workbook(io.BytesIO(resp.content))


def test_all():
    tmp = tempfile.mkdtemp(prefix="amr_api_test_")
    try:
        now, (r1, r2, r3) = build(tmp)
        app = FastAPI()
        app.include_router(api.router)
        # the real app mounts static files on "/" AFTER including routers
        web = os.path.join(tmp, "web"); os.makedirs(web)
        open(os.path.join(web, "index.html"), "w").write("<html>ok</html>")
        app.mount("/", StaticFiles(directory=web, html=True), name="static")
        c = TestClient(app)

        print("[meta]")
        m = c.get("/api/history/meta").json()
        expect(m["status"] == "ok" and m["enabled"] is True, "meta enabled")
        expect({"NAVIGATION", "SLAM", "SYSTEM"} <= set(m["categories"]), f"categories {m['categories']}")
        expect(set(m["run_types"]) == {"goal", "waypoints"}, f"run_types {m['run_types']}")
        expect({"10.0.0.5", "10.0.0.6"} <= set(m["clients"]), f"clients {m['clients']}")
        expect(m["runs_first_ts"] is not None and m["excel_available"] is True, "runs range + excel_available")
        expect("/api/history" in str(c.get("/api/history/meta").url) and
               c.get("/index.html").text == "<html>ok</html>",
               "router answers AND static mount still serves files (router included before the mount)")

        print("[events]")
        r = c.get("/api/history/events").json()
        expect(r["enabled"] and r["total"] >= 4, f"events total {r['total']} (backend_started + 3 logged)")
        by = {e["action"]: e for e in r["items"]}
        expect(by["start_slam"]["result"] == "success" and "SLAM started" in by["start_slam"]["detail"],
               "start_slam row has result=success and message appended")
        expect(by["nav_goal"]["result"] == "failed", "nav_goal row has result=failed")
        expect(c.get("/api/history/events", params={"category": "SLAM"}).json()["total"] == 1, "filter category=SLAM")
        expect(c.get("/api/history/events", params={"result": "failed"}).json()["total"] == 1, "filter result=failed")
        expect(c.get("/api/history/events", params={"q": "lidar"}).json()["total"] == 1, "text search q=lidar")
        expect(c.get("/api/history/events", params={"from": now + 1000}).json()["total"] == 0, "from in the future -> 0")
        expect(c.get("/api/history/events", params={"to": now - 1000}).json()["total"] == 0, "to in the past -> 0")
        pg = c.get("/api/history/events", params={"page_size": 2, "page": 1}).json()
        expect(len(pg["items"]) == 2 and pg["total"] >= 4, "paging page_size=2")
        expect(c.get("/api/history/events", params={"page_size": 500}).status_code == 422, "page_size > 200 rejected")

        print("[runs]")
        r = c.get("/api/history/runs").json()
        expect(r["total"] == 3, f"3 runs (got {r['total']})")
        expect([x["id"] for x in r["items"]] == [r3, r2, r1], "newest first")
        expect(c.get("/api/history/runs", params={"result": "SUCCEEDED"}).json()["total"] == 2, "filter result=SUCCEEDED")
        expect(c.get("/api/history/runs", params={"type": "waypoints"}).json()["total"] == 1, "filter type=waypoints")
        expect(c.get("/api/history/runs", params={"scenario": "A to"}).json()["total"] == 2, "filter scenario (LIKE)")
        expect(c.get("/api/history/runs", params={"map_id": "map_01"}).json()["total"] == 3, "filter map_id")
        d = c.get(f"/api/history/runs/{r1}").json()["run"]
        expect(d["result"] == "SUCCEEDED" and abs(d["path_length_m"] - 1.0) < 0.05, f"detail path_length {d['path_length_m']}")
        expect(abs(d["pos_error_m"] - math.hypot(0.002, 0.004)) < 1e-6, "detail pos_error")
        expect(len(d["trajectory"]) >= 11 and d["waypoints"] == [], "detail has trajectory, no waypoints")
        d3 = c.get(f"/api/history/runs/{r3}").json()["run"]
        expect([w["name"] for w in d3["waypoints"]] == ["P1", "P2"], "waypoint run detail lists P1, P2")
        expect(c.get("/api/history/runs/9999").json()["status"] == "error", "unknown run -> error json")

        print("[patch]")
        p = c.patch(f"/api/history/runs/{r1}", json={"scenario": "  Straight 1 m ", "notes": "tape ok",
                                                      "measured_error_cm": "1.5"}).json()
        expect(p["status"] == "ok" and p["scenario"] == "Straight 1 m", "patch ok, scenario trimmed")
        flush()
        d = c.get(f"/api/history/runs/{r1}").json()["run"]
        expect(d["scenario"] == "Straight 1 m" and d["notes"] == "tape ok" and d["measured_error_cm"] == 1.5,
               "patched values are stored")
        expect(c.patch(f"/api/history/runs/{r1}", json={"measured_error_cm": "abc"}).json()["status"] == "error", "bad number rejected")
        expect(c.patch(f"/api/history/runs/{r1}", json={"measured_error_cm": -3}).json()["status"] == "error", "negative rejected")
        expect(c.patch(f"/api/history/runs/{r1}", json={}).json()["status"] == "error", "empty patch rejected")
        expect(c.patch("/api/history/runs/9999", json={"notes": "x"}).json()["status"] == "error", "patch unknown run rejected")
        c.patch(f"/api/history/runs/{r1}", json={"notes": ""}); flush()
        expect(c.get(f"/api/history/runs/{r1}").json()["run"]["notes"] == "", "empty string clears notes")

        print("[export runs.xlsx]")
        resp = c.get("/api/history/export/runs.xlsx")
        expect(resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"], "200 + xlsx content-type")
        expect("attachment" in resp.headers["content-disposition"] and ".xlsx" in resp.headers["content-disposition"],
               f"download filename: {resp.headers['content-disposition']}")
        wb = xl(resp)
        expect(wb.sheetnames == ["Runs", "Waypoints", "Trajectory", "Summary"], f"sheets {wb.sheetnames}")
        ws = wb["Runs"]
        hdr = [c_.value for c_ in ws[1]]
        expect(hdr[:3] == ["Run ID", "Date", "Start time"], f"header starts {hdr[:3]}")
        expect({"Start X (m)", "Goal Yaw (deg)", "End Yaw (deg)", "Position error (m)", "Yaw error (deg)",
                "Measured error (cm)", "Path length (m)"} <= set(hdr), "required columns present")
        expect(ws.max_row == 4, f"3 data rows + header (max_row={ws.max_row})")
        col = {h: i + 1 for i, h in enumerate(hdr)}
        first = ws.cell(row=2, column=col["Run ID"]).value
        expect(first == r1, "chronological order (oldest first)")
        for h in ("Start X (m)", "Goal Yaw (deg)", "End Yaw (deg)", "Position error (m)", "Yaw error (deg)", "Path length (m)"):
            expect(ws.cell(row=2, column=col[h]).number_format == "0.000", f"'{h}' formatted 0.000")
        expect(abs(ws.cell(row=2, column=col["End X (m)"]).value - 1.002) < 1e-9, "full precision kept in the cell")
        expect(ws.cell(row=2, column=col["Date"]).number_format == "yyyy-mm-dd", "date is a real date cell")
        expect(ws.cell(row=2, column=col["Measured error (cm)"]).value == 1.5, "measured error exported")
        sm = wb["Summary"]
        flat = {(r_[0].value): [c_.value for c_ in r_] for r_ in sm.iter_rows() if r_[0].value}
        expect("ALL RUNS" in flat, "summary has ALL RUNS row")
        allrow = flat["ALL RUNS"]
        expect(allrow[1] == 3 and allrow[2] == 2 and abs(allrow[3] - 66.6667) < 0.01, f"ALL RUNS: runs/succeeded/rate {allrow[1:4]}")
        expect("A to B" in flat and flat["A to B"][1] == 1 or "Straight 1 m" in flat, "per-scenario rows present")
        expect(len(wb["Waypoints"]["A"]) and wb["Waypoints"].max_row == 3, "waypoint sheet has the 2 waypoint rows")
        expect(wb["Trajectory"].max_row > 20, "trajectory rows exported")

        resp = c.get("/api/history/export/runs.xlsx", params={"result": "SUCCEEDED", "include_trajectory": "false"})
        wb = xl(resp)
        expect(wb["Runs"].max_row == 3 and "Trajectory" not in wb.sheetnames, "filters respected; trajectory optional")

        resp = c.get("/api/history/export/runs.xlsx", params={"run_id": r1})
        wb = xl(resp)
        expect(wb["Runs"].max_row == 2 and len(wb["Trajectory"]._charts) == 1, "single-run export has 1 row + XY chart")
        expect(f"amr_run_{r1}_" in resp.headers["content-disposition"], "single-run file name")

        print("[export events.xlsx]")
        resp = c.get("/api/history/export/events.xlsx", params={"category": "SLAM"})
        wb = xl(resp)
        ws = wb["Events"]
        expect(ws.max_row == 2 and ws.cell(row=2, column=6).value == "start_slam", "events export filtered (1 row)")
        expect(ws.cell(row=2, column=2).number_format == "yyyy-mm-dd", "events date cell")

        print("[disabled]")
        db.shutdown()
        j = c.get("/api/history/events").json()
        expect(j["enabled"] is False and j["items"] == [], "disabled -> empty page, no error")
        expect(c.get("/api/history/export/runs.xlsx").status_code == 503, "disabled export -> 503")
        expect(c.get("/api/history/meta").json()["enabled"] is False, "disabled meta")
    finally:
        db.shutdown()
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    try:
        test_all()
    except Exception:
        import traceback
        traceback.print_exc()
        _FAILURES.append("test_all raised")
    print()
    if _FAILURES:
        print(f"{len(_FAILURES)} FAILURE(S):")
        for f in _FAILURES:
            print("  -", f)
        sys.exit(1)
    print("All tests passed.")